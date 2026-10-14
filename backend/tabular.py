"""Bounded offline readers. Imported files are data; no scripts or macros run."""
import csv
from datetime import date, datetime, time
from email import policy
from email.parser import BytesParser
from io import BytesIO, StringIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile, BadZipFile

MAX_BYTES = 10_000_000
MAX_ROWS = 20_000
MAX_COLUMNS = 200
MAX_CELLS = 300_000


def cell_text(value):
    if value is None: return ''
    if isinstance(value, bool): return 'true' if value else 'false'
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time() else value.isoformat()
    if isinstance(value, time): return value.isoformat(timespec='minutes' if not value.second and not value.microsecond else 'seconds')
    if isinstance(value, date): return value.isoformat()
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    result = str(value).strip()
    if len(result) > 65536:
        raise ValueError('Ячейка превышает предел 65 536 символов.')
    return result


def decoded(raw, encoding='auto'):
    if encoding != 'auto': return raw.decode(encoding)
    for codec in ['utf-8-sig', 'utf-16' if raw[:2] in [b'\xff\xfe', b'\xfe\xff'] else 'cp1251']:
        try: return raw.decode(codec)
        except UnicodeError: pass
    raise ValueError('Не удалось определить кодировку. Выберите её вручную.')


def bounded(sheets):
    rows = cells = 0
    for sheet in sheets:
        rows += len(sheet['rows'])
        if rows > MAX_ROWS: raise ValueError('В файле больше 20 000 строк. Разделите его на части.')
        for row in sheet['rows']:
            if len(row) > MAX_COLUMNS: raise ValueError('В таблице больше 200 колонок.')
            cells += len(row)
            if cells > MAX_CELLS: raise ValueError('В файле больше 300 000 ячеек.')
        sheet['rows'] = [[cell_text(v) for v in row] for row in sheet['rows']]
    if not sheets: raise ValueError('Таблицы в файле не найдены.')
    return sheets


def xml_root(raw):
    text = decoded(raw, 'utf-8-sig')
    if '\x00' in text or '<!DOCTYPE' in text.upper() or '<!ENTITY' in text.upper():
        raise ValueError('XML с DTD или сущностями не поддерживается.')
    return ET.fromstring(text)


def flat_xml(raw):
    root = xml_root(raw)
    if root.tag == 'atlas':
        raise ValueError('Это полная XML-копия Атласа. Используйте восстановление XML ниже на странице.')
    containers = [c for c in root if len(c) and all(g.attrib or len(g) for g in c)] or [root]
    sheets = []
    for container in containers:
        records = []
        for item in container:
            if len(records) >= MAX_ROWS: raise ValueError('Слишком много записей XML.')
            record = dict(item.attrib)
            for child in item:
                if len(child) or child.attrib:
                    raise ValueError('Табличный XML должен содержать плоские записи: атрибуты или простые дочерние поля.')
                if child.tag in record: raise ValueError('Повторяющееся поле XML.')
                record[child.tag] = child.text or ''
            if record: records.append(record)
        if records:
            headers = list(dict.fromkeys(k for r in records for k in r))
            sheets.append(dict(name=container.tag, rows=[headers]+[[r.get(h, '') for h in headers] for r in records], formulas=[]))
    return sheets


def html_tables(raw, encoding):
    from bs4 import BeautifulSoup
    message = BytesParser(policy=policy.default).parsebytes(raw)
    html_parts = []
    for part in message.walk():
        if part.get_content_type() == 'text/html':
            payload = part.get_payload(decode=True)
            if payload:
                charset = part.get_content_charset()
                html_parts.append(payload.decode(charset) if charset else decoded(payload, encoding))
    if not html_parts: html_parts = [decoded(raw, encoding)]
    result = []
    for html in html_parts:
        soup = BeautifulSoup(html, 'html.parser')
        for hidden in soup(['script', 'style']): hidden.decompose()
        for table in soup.find_all('table'):
            grid, row_count = {}, 0
            for ri, tr in enumerate(table.find_all('tr')):
                if tr.find_parent('table') is not table: continue
                if ri >= MAX_ROWS: raise ValueError('Слишком много строк HTML.')
                column = 0
                for td in tr.find_all(['td', 'th'], recursive=False):
                    while (ri, column) in grid: column += 1
                    span, down = int(td.get('colspan', 1)), int(td.get('rowspan', 1))
                    if span < 1 or down < 1 or column+span > MAX_COLUMNS or ri+down > MAX_ROWS or len(grid)+span*down > MAX_CELLS:
                        raise ValueError('Объединённая ячейка выходит за допустимый размер таблицы.')
                    for br in td.find_all('br'): br.replace_with('\n')
                    value = td.get_text().strip()
                    for r in range(ri, ri+down):
                        for c in range(column, column+span): grid[r,c] = value
                    column += span
                    row_count = max(row_count, ri+down)
            if grid:
                columns = max(c for _, c in grid)+1
                if row_count*columns > MAX_CELLS: raise ValueError('Слишком большая HTML-таблица.')
                result.append(dict(name=f'Таблица {len(result)+1}', rows=[[grid.get((r,c),'') for c in range(columns)] for r in range(row_count)], formulas=[]))
    return result


def parse_file(filename, raw, delimiter='auto', encoding='auto'):
    if not raw or len(raw) > MAX_BYTES: raise ValueError('Файл пустой или превышает 10 МБ.')
    ext = Path(filename).suffix.lower()
    warnings = []
    try:
        if ext in ['.xlsx', '.xlsm']:
            from openpyxl import load_workbook
            with ZipFile(BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 50_000_000 or len(archive.infolist()) > 2000:
                    raise ValueError('Книга превышает допустимый распакованный размер.')
                for item in archive.infolist():
                    if item.filename.lower().endswith('.xml'):
                        text = archive.read(item).replace(b'\x00',b'').upper()
                        if b'<!DOCTYPE' in text or b'<!ENTITY' in text:
                            raise ValueError('Объявления сущностей в книге не поддерживаются.')
            wb = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            sheets = []
            total_cells = 0
            try:
                for ws in wb.worksheets:
                    if (ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > MAX_COLUMNS:
                        raise ValueError('Лист превышает 20 000 строк или 200 колонок. Удалите лишнее форматирование/пустые области.')
                    rows, formulas = [], []
                    for ri, row in enumerate(ws.iter_rows()):
                        total_cells += len(row)
                        if ri >= MAX_ROWS or total_cells > MAX_CELLS: raise ValueError('Книга превышает предел строк/ячеек.')
                        rows.append([c.value for c in row])
                        formulas += [(ri, ci) for ci, c in enumerate(row) if c.data_type == 'f']
                    sheets.append(dict(name=ws.title, rows=rows, formulas=formulas))
            finally: wb.close()
            if any(s['formulas'] for s in sheets): warnings.append('Есть формулы. Сопоставленные ячейки с формулами будут отклонены: сохраните их как значения.')
            if ext == '.xlsm': warnings.append('Макросы не запускаются и не переносятся.')
        elif ext == '.xls' and raw.startswith((bytes.fromhex('D0CF11E0A1B11AE1'), b'\x09\x08', b'\x09\x04', b'\x09\x02', b'\x09\x00')):
            import xlrd
            wb = xlrd.open_workbook(file_contents=raw, on_demand=True)
            sheets = []
            try:
                for ws in wb.sheets():
                    if ws.nrows > MAX_ROWS or ws.ncols > MAX_COLUMNS or ws.nrows*ws.ncols > MAX_CELLS:
                        raise ValueError('Лист XLS превышает допустимый размер.')
                    rows = []
                    for ri in range(ws.nrows):
                        row = []
                        for cell in ws.row(ri):
                            if cell.ctype == xlrd.XL_CELL_ERROR: raise ValueError(f'Ошибка Excel в строке {ri+1}.')
                            value = xlrd.xldate_as_datetime(cell.value, wb.datemode) if cell.ctype == xlrd.XL_CELL_DATE else bool(cell.value) if cell.ctype == xlrd.XL_CELL_BOOLEAN else cell.value
                            if cell.ctype == xlrd.XL_CELL_DATE and 0 <= cell.value < 1: value = value.time()
                            row.append(value)
                        rows.append(row)
                    sheets.append(dict(name=ws.name, rows=rows, formulas=[]))
            finally: wb.release_resources()
            warnings.append('Старый XLS читается по сохранённым значениям; формулы не пересчитываются.')
        elif ext in ['.html', '.htm', '.mhtml', '.mht', '.xls']:
            sheets = html_tables(raw, encoding)
            warnings.append('Объединённые ячейки развёрнуты. Проверьте строку заголовков и первые строки; скрипты и внешние ресурсы не выполняются.')
        elif ext == '.xml':
            sheets = flat_xml(raw)
            warnings.append('Плоские записи XML. Вложенные форматы других систем требуют отдельного преобразования.')
        elif ext in ['.csv', '.tsv', '.txt']:
            text = decoded(raw, encoding)
            separator = '\t' if ext == '.tsv' and delimiter == 'auto' else delimiter
            if separator == 'auto':
                try: separator = csv.Sniffer().sniff(text[:8192], delimiters=';,\t|').delimiter
                except csv.Error: separator = ';'
            rows = []
            total_cells = 0
            for row in csv.reader(StringIO(text), delimiter=separator):
                if len(rows) >= MAX_ROWS: raise ValueError('В файле больше 20 000 строк.')
                total_cells += len(row)
                if len(row) > MAX_COLUMNS or total_cells > MAX_CELLS: raise ValueError('Файл превышает предел колонок/ячеек.')
                rows.append(row)
            sheets = [dict(name=Path(filename).stem, rows=rows, formulas=[])]
        else:
            raise ValueError('Поддерживаются CSV, TSV, XLSX, XLS, HTML, MHTML и плоский XML. Для .roz нужен экспорт из aSc.')
        return dict(sheets=bounded(sheets), warnings=warnings)
    except (UnicodeError, LookupError, BadZipFile, ET.ParseError, csv.Error, OSError, OverflowError) as exc:
        raise ValueError('Не удалось прочитать файл: ' + str(exc)) from exc
