"""Offline workbook/XML exchange and escaped, standalone print documents."""
from collections import defaultdict
from html import escape
from io import BytesIO
import json
from typing import Literal
from xml.etree import ElementTree as ET

from pydantic import Field
from .models import Model, State
from .resources import teacher_ids, room_ids


class ExportRequest(Model):
    revision: int
    format: Literal['xlsx', 'xlsx_data', 'html', 'xml'] = 'html'
    view: Literal['class', 'teacher', 'room', 'group', 'student'] = 'class'
    entityIds: list[str] = Field(default_factory=list, max_length=5000)
    cycleIds: list[str] = Field(default_factory=list, max_length=32)
    template: Literal['grid', 'list', 'overview'] = 'grid'
    paper: Literal['A4', 'A3', 'A2', 'A1'] = 'A4'
    landscape: bool = True
    showTeachers: bool = True
    showRooms: bool = True
    showTimes: bool = True
    title: str = Field(default='', max_length=160)


TABLES = {'classes': 'Классы', 'teachers': 'Учителя', 'rooms': 'Кабинеты', 'subjects': 'Предметы',
          'students': 'Ученики', 'groups': 'Группы', 'profiles': 'Звонки', 'workloads': 'Нагрузка',
          'placements': 'Занятия', 'events': 'Календарь', 'rules': 'Ограничения', 'courses': 'Курсы', 'courseRequests': 'Заявки',
          'datedLessons': 'Изменения по датам', 'absences': 'Отсутствия учителей', 'commitments': 'Бронирования и дежурства'}
VIEWS = dict(class_='classes', teacher='teachers', room='rooms', group='groups', student='students')


def schedule_rows(state, options):
    table = VIEWS['class_' if options.view == 'class' else options.view]
    entities = {x.id: x for x in getattr(state, table)}
    if set(options.entityIds) - entities.keys() or set(options.cycleIds) - {c.id for c in state.cycles.patterns}:
        raise ValueError('Выбраны отсутствующие записи или недели.')
    chosen = [x for x in entities.values() if not options.entityIds or x.id in options.entityIds]
    if options.view == 'student' and not options.entityIds:
        chosen = [x for x in chosen if x.active]
    cycles = [c for c in state.cycles.patterns if not options.cycleIds or c.id in options.cycleIds]
    groups = {g.id: g for g in state.groups}
    loads = {w.id: w for w in state.workloads}
    profiles = {p.id: p for p in state.profiles}
    names = {key: {x.id: x.name for x in getattr(state, key)} for key in TABLES if key not in ['placements', 'workloads', 'courseRequests']}
    pupil_classes = {x.id: x.classId for x in state.students}
    chosen_ids = {e.id for e in chosen}
    cycle_ids = {c.id for c in cycles}
    rows = []
    for p in state.placements:
        if p.cycleId not in cycle_ids:
            continue
        w = loads[p.workloadId]
        g = groups[w.groupId]
        if options.view == 'class':
            ids = set(g.classIds) if g.kind == 'whole' or g.rosterStatus in ['unknown', 'partial'] or g.rosterStatus is None and not g.studentIds else {pupil_classes[id] for id in g.studentIds}
        elif options.view == 'teacher':
            ids = set(teacher_ids(w))
        elif options.view == 'room':
            ids = set(room_ids(w, p.roomId))
        elif options.view == 'group':
            ids = {g.id}
        else:
            ids = {id for id, cls in pupil_classes.items() if cls in g.classIds} if g.kind == 'whole' else set(g.studentIds)
        slot = profiles[w.profileId].for_day(p.day)[p.period]
        for id in sorted(ids & chosen_ids):
            rows.append(dict(entityId=id, entity=entities[id].name, cycleId=p.cycleId,
                             cycle=next(c.name for c in cycles if c.id == p.cycleId), day=p.day, dayName=state.settings.days[p.day],
                             period=p.period, label=slot.label, start=slot.start, end=slot.end,
                             subject=names['subjects'][w.subjectId], group=g.name,
                             teachers=', '.join(names['teachers'][id] for id in teacher_ids(w)) or ('Не указан' if w.teacherRequired else 'Без учителя'),
                             rooms=', '.join(names['rooms'][id] for id in room_ids(w, p.roomId)) or ('Не указан' if w.roomRequired else 'Без кабинета'),
                             placementId=p.id, locked=p.locked))
    order = {e.id: i for i, e in enumerate(chosen)}
    cycle_order = {c.id: i for i, c in enumerate(cycles)}
    rows.sort(key=lambda r: (order[r['entityId']], cycle_order[r['cycleId']], r['day'], r['start'], r['subject'], r['placementId']))
    return chosen, cycles, rows


def state_xml(state):
    root = ET.Element('atlas', version='1')

    def put(node, value):
        if isinstance(value, dict):
            node.set('type', 'object')
            for key, item in value.items():
                put(ET.SubElement(node, 'field', name=key), item)
        elif isinstance(value, list):
            node.set('type', 'array')
            for item in value:
                put(ET.SubElement(node, 'item'), item)
        elif value is None:
            node.set('type', 'null')
        else:
            node.set('type', 'bool' if isinstance(value, bool) else 'int' if isinstance(value, int) else 'string')
            node.text = str(value).lower() if isinstance(value, bool) else str(value)
    put(ET.SubElement(root, 'state'), state.model_dump(mode='json'))
    return ET.tostring(root, encoding='utf-8', xml_declaration=True)


def read_state_xml(raw):
    if len(raw) > 20_000_000:
        raise ValueError('XML превышает допустимый размер.')
    text = raw.decode('utf-8-sig')
    if '\x00' in text or '<!DOCTYPE' in text.upper() or '<!ENTITY' in text.upper():
        raise ValueError('XML превышает допустимый размер или содержит запрещённое объявление сущностей.')
    root = ET.fromstring(text)
    if root.tag != 'atlas' or root.attrib != {'version': '1'} or len(root) != 1 or root[0].tag != 'state':
        raise ValueError('Нужна XML-копия Атласа версии 1. XML других систем требует сопоставления их формата.')

    def get(node, depth=0):
        if depth > 30:
            raise ValueError('Слишком глубокая структура XML.')
        kind = node.get('type')
        if kind == 'object':
            keys = [c.get('name') for c in node]
            if any(c.tag != 'field' for c in node) or None in keys or len(set(keys)) != len(keys):
                raise ValueError('Повторяющиеся или неверные поля XML.')
            return {c.get('name'): get(c, depth+1) for c in node}
        if kind == 'array':
            if any(c.tag != 'item' for c in node):
                raise ValueError('Неверный список XML.')
            return [get(c, depth+1) for c in node]
        if len(node):
            raise ValueError('Неверное простое значение XML.')
        text = node.text or ''
        if kind == 'string': return text
        if kind == 'null': return None
        if kind == 'int': return int(text)
        if kind == 'bool' and text in ['true', 'false']: return text == 'true'
        raise ValueError('Неизвестный тип XML.')
    return State.model_validate(get(root[0]))


def workbook(state, options):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row_index, row in enumerate(rows, 2):
            values = [json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v for v in row]
            if any(isinstance(v, str) and len(v) > 32767 for v in values):
                raise ValueError('Поле превышает предел Excel 32767 символов. Для полной копии используйте XML или JSON.')
            if any(isinstance(v, str) and ILLEGAL_CHARACTERS_RE.search(v) for v in values):
                raise ValueError('В данных есть управляющие символы, недопустимые в Excel. Исправьте текст или используйте JSON-копию.')
            ws.append(values)
            for column, value in enumerate(values, 1):
                if isinstance(value, str):
                    ws.cell(row_index, column).data_type = 's'  # User text must never become an Excel formula.
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='28674F')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for i in range(1, len(headers)+1):
            values = [str(ws.cell(row, i).value or '') for row in range(1, min(ws.max_row, 150)+1)]
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(14, max(map(len, values))+2))
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        return ws

    if options.format == 'xlsx_data':
        sheet('О формате', ['Параметр', 'Значение'], [('Формат', 'Таблицы Атласа 1'), ('Версия базы', state.revision),
              ('Назначение', 'Данные всей базы. Заголовки — имена полей для обмена; списки и объекты записаны JSON.'),
              ('Полная копия', 'Для восстановления всей базы используйте XML или JSON. Пустые таблицы не удаляют данные при импорте.')])
        data = state.model_dump(mode='json')
        for key, name in TABLES.items():
            model = State.model_fields[key].annotation.__args__[0]
            headers = list(model.model_fields)
            sheet(name, headers, [[item[h] for h in headers] for item in data[key]])
        sheet('Настройки', ['Поле', 'Значение'], [(k, data[k]) for k in ['schemaVersion', 'settings', 'cycles', 'campus', 'importInfo']])
    else:
        _, _, rows = schedule_rows(state, options)
        columns = [('entity', 'Для кого'), ('cycle', 'Неделя'), ('dayName', 'День'), ('label', 'Урок')]
        if options.showTimes: columns += [('start', 'Начало'), ('end', 'Конец')]
        columns += [('subject', 'Предмет'), ('group', 'Группа')]
        if options.showTeachers: columns += [('teachers', 'Учителя')]
        if options.showRooms: columns += [('rooms', 'Кабинеты')]
        sheet('Расписание', [label for _, label in columns], [[row[key] for key, _ in columns] for row in rows])
        sheet('О расписании', ['Параметр', 'Значение'], [('Школа', state.settings.name), ('Версия', state.revision), ('Период', f'{state.settings.yearStart} — {state.settings.yearEnd}'), ('Основа', 'Повторяющиеся шаблоны недель; изменения по датам сюда не входят.')])
    out = BytesIO(); wb.save(out)
    return out.getvalue()


def html_document(state, options):
    entities, cycles, rows = schedule_rows(state, options)
    esc = lambda value: escape(str(value), quote=True)
    title = options.title or state.settings.name
    by_entity = defaultdict(list)
    by_cell = defaultdict(list)
    for row in rows:
        by_entity[row['entityId'], row['cycleId']].append(row)
        by_cell[row['entityId'], row['cycleId'], row['day'], row['period']].append(row)
    periods = max((len(p.periods) for p in state.profiles), default=1)
    if len(entities) * len(cycles) * len(state.settings.days) * periods > 250_000:
        raise ValueError('Слишком большой пакет печати. Выберите меньше записей или недель.')

    def card(r):
        details = [r['group']]
        if options.showTimes: details += [f'{r["start"]}–{r["end"]}']
        if options.showTeachers: details += [r['teachers']]
        if options.showRooms: details += [r['rooms']]
        return '<div class="lesson"><strong>' + esc(r['subject']) + '</strong><small>' + '<br>'.join(esc(x) for x in details) + '</small></div>'

    def header(name, cycle):
        return '<h1>' + esc(title) + '</h1><h2>' + esc(name) + ' · ' + esc(cycle.name) + '</h2><p class="meta">' + esc(f'{state.settings.yearStart} — {state.settings.yearEnd} · версия {state.revision}') + '</p>'
    body = []
    if options.template == 'overview':
        for c in cycles:
            content = header('Сводное расписание', c) + '<table class="overview"><thead><tr><th rowspan="2">Для кого</th>'
            content += ''.join(f'<th colspan="{periods}">{esc(day)}</th>' for day in state.settings.days) + '</tr><tr>'
            content += ''.join(f'<th>{p}</th>' for _ in state.settings.days for p in range(periods)) + '</tr></thead><tbody>'
            for e in entities:
                content += '<tr><th>' + esc(e.name) + '</th>'
                content += ''.join('<td>' + ''.join(card(r) for r in by_cell[e.id,c.id,d,p]) + '</td>' for d in range(len(state.settings.days)) for p in range(periods))
                content += '</tr>'
            body.append('<section>' + content + '</tbody></table></section>')
    else:
        for e in entities:
            for c in cycles:
                content = header(e.name, c)
                if options.template == 'grid':
                    content += '<table><thead><tr><th>Урок</th>' + ''.join('<th>'+esc(d)+'</th>' for d in state.settings.days) + '</tr></thead><tbody>'
                    for p in range(periods):
                        content += f'<tr><th>{p}</th>' + ''.join('<td>'+''.join(card(r) for r in by_cell[e.id,c.id,d,p])+'</td>' for d in range(len(state.settings.days))) + '</tr>'
                    content += '</tbody></table>'
                else:
                    content += '<table><thead><tr><th>День</th><th>Урок</th><th>Занятие</th></tr></thead><tbody>'
                    for r in by_entity[e.id,c.id]:
                        content += '<tr><td>'+esc(r['dayName'])+'</td><td>'+esc(r['label'])+'</td><td>'+card(r)+'</td></tr>'
                    content += '</tbody></table>'
                if not by_entity[e.id,c.id]: content += '<p>Нет размещённых занятий.</p>'
                body.append('<section>'+content+'</section>')
    size = dict(A4='297mm 210mm', A3='420mm 297mm', A2='594mm 420mm', A1='841mm 594mm')[options.paper].split()
    if not options.landscape: size.reverse()
    return '<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'+esc(title)+'</title><style>' + '''
body{font:12px Arial,sans-serif;color:#243b30;background:#fff;margin:20px}h1{font-size:19px;margin:0 0 5px}h2{font-size:16px;margin:5px 0}.meta{font-size:10px;color:#657567}section{margin:0 0 30px;break-after:page}section:last-child{break-after:auto}table{border-collapse:collapse;width:100%;table-layout:fixed}th,td{border:1px solid #acb9ae;padding:5px;vertical-align:top;overflow-wrap:anywhere}th{background:#edf2ed;font-size:10px}th:first-child{width:42px}.lesson{padding:2px 0}.lesson+.lesson{border-top:1px dashed #bcc6be;margin-top:4px}small{display:block;font-size:9px;line-height:1.3}.overview{font-size:7px}.overview small{font-size:6px}.overview th:first-child{width:70px}.overview td{padding:2px}thead{display:table-header-group}tr{break-inside:avoid}
''' + '@page{size:'+' '.join(size)+';margin:8mm}@media print{body{margin:0}section{margin-bottom:0}}' + '</style></head><body>' + ''.join(body) + '</body></html>'


class XmlRestore(Model):
    revision: int
    xml: str = Field(max_length=20_000_000)


def register_exports(app, Workspace, require_revision, create_preview):
    from fastapi import HTTPException, Response

    @app.post('/api/exchange/export')
    def export(body: ExportRequest, workspace: Workspace = 'school'):
        state = require_revision(workspace, body.revision)
        try:
            if body.format == 'xml':
                content, ext, media = state_xml(state), 'xml', 'application/xml'
            elif body.format in ['xlsx', 'xlsx_data']:
                content, ext, media = workbook(state, body), 'xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            else:
                content, ext, media = html_document(state, body), 'html', 'text/html'
        except ValueError as exc:
            raise HTTPException(422, str(exc))
        return Response(content=content, media_type=media, headers={'Content-Disposition': f'attachment; filename="atlas-{body.format}.{ext}"', 'Cache-Control': 'no-store'})

    @app.post('/api/exchange/xml-preview')
    def preview_xml(body: XmlRestore, workspace: Workspace = 'school'):
        state = require_revision(workspace, body.revision)
        try:
            candidate = read_state_xml(body.xml.encode('utf-8'))
            candidate.revision = state.revision
        except (ValueError, ET.ParseError) as exc:
            raise HTTPException(422, str(exc))
        return create_preview(workspace, state, candidate, allow_issues=True)
