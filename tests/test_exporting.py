from io import BytesIO
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend import api
from backend.models import State, Placement, Cycles, Course, CourseRequest
from backend.seed import initial_state
from backend.exporting import ExportRequest, schedule_rows, html_document, workbook, state_xml, read_state_xml


def fixture():
    s = initial_state('demo')
    s.placements = [Placement(id='a', workloadId='w0', day=0, period=1, roomId='r0'),
                    Placement(id='m', workloadId='w4', day=1, period=2, roomId='r1'),
                    Placement(id='f', workloadId='w5', day=1, period=2, roomId='r2')]
    return s


def test_selected_people_only_receive_their_groups_and_real_clock():
    s = fixture()
    s.profiles[1].dayPeriods = {'0': [p.model_copy() for p in s.profiles[1].periods]}
    s.profiles[1].dayPeriods['0'][1].start = '08:30'
    _, _, rows = schedule_rows(s, ExportRequest(revision=0, view='student', entityIds=['s1']))
    assert {r['placementId'] for r in rows} == {'a', 'm'}
    assert rows[0]['start'] == '08:30'
    _, _, rows = schedule_rows(s, ExportRequest(revision=0, view='student', entityIds=['s2']))
    assert {r['placementId'] for r in rows} == {'a', 'f'}
    with pytest.raises(ValueError): schedule_rows(s, ExportRequest(revision=0, entityIds=['missing']))


def test_coteachers_rooms_and_week_filters_are_preserved():
    s = fixture()
    s.workloads[0].additionalTeacherIds = ['t0']
    s.workloads[0].roomIds = ['r0']
    s.workloads[0].requiredRoomIds = ['r2']
    s.cycles = Cycles(patterns=[dict(id='weekly', name='A'), dict(id='b', name='B')])
    s.placements[0].cycleId = 'b'
    _, _, rows = schedule_rows(s, ExportRequest(revision=0, view='teacher', entityIds=['t0'], cycleIds=['b']))
    assert len(rows) == 1 and 'Учитель английского' in rows[0]['teachers'] and 'Учитель математики' in rows[0]['teachers']
    _, _, rows = schedule_rows(s, ExportRequest(revision=0, view='room', entityIds=['r2'], cycleIds=['b']))
    assert len(rows) == 1 and rows[0]['rooms'] == '301, 303'


def test_standalone_html_escapes_names_and_hides_selected_details():
    s = fixture()
    s.subjects[2].name = '<img src=x onerror=alert(1)>'
    doc = html_document(s, ExportRequest(revision=0, entityIds=['c-demo-a'], showTeachers=False, showRooms=False, paper='A3', landscape=False))
    assert '<img' not in doc and '&lt;img' in doc and 'Учитель английского' not in doc
    assert '297mm 420mm' in doc and doc.count('<section>') == 1
    for template in ['list', 'overview']:
        doc = html_document(s, ExportRequest(revision=0, template=template))
        assert '10A' in doc and '10D' in doc and '<script' not in doc


def test_xml_roundtrip_includes_cycles_and_requests_and_refuses_entities():
    s = fixture()
    s.courses = [Course(id='m', name='Math & core', sections=[dict(groupId='g-math')])]
    s.courseRequests = [CourseRequest(id='choice', studentId='s1', courseIds=['m'])]
    assert read_state_xml(state_xml(s)) == s
    for bad in [b'<!DOCTYPE atlas [<!ENTITY x "boom">]><atlas version="1"><state>&x;</state></atlas>',
                b'<another/>', b'<atlas version="1"><state type="object"><field name="revision" type="int">0</field><field name="revision" type="int">1</field></state></atlas>']:
        with pytest.raises(ValueError): read_state_xml(bad)


def test_xlsx_does_not_execute_text_as_formulas_and_data_columns_are_lossless():
    s = fixture(); s.subjects[2].name = '=HYPERLINK("https://example.invalid")'
    output = workbook(s, ExportRequest(revision=0, format='xlsx', entityIds=['c-demo-a']))
    wb = load_workbook(BytesIO(output), data_only=False)
    ws = wb['Расписание']
    assert ws['G2'].value == s.subjects[2].name and ws['G2'].data_type == 's'
    assert ws.freeze_panes == 'A2' and ws.auto_filter.ref
    data = load_workbook(BytesIO(workbook(s, ExportRequest(revision=0, format='xlsx_data'))))
    headers = [c.value for c in data['Нагрузка'][1]]
    assert 'cycleHours' in headers and 'additionalTeacherIds' in headers
    assert data['Занятия'].max_row == len(s.placements)+1


def test_long_excel_field_is_rejected_without_truncation():
    s = fixture(); s.groups[0].rosterNote = 'x'*32768
    with pytest.raises(ValueError, match='32767'): workbook(s, ExportRequest(revision=0, format='xlsx_data'))


def test_export_and_xml_preview_use_revision_and_never_implicitly_save(tmp_path, monkeypatch):
    monkeypatch.setattr(api, 'DB', tmp_path / 'export.db')
    client = TestClient(api.app)
    original = client.get('/api/state?workspace=demo').json()
    response = client.post('/api/exchange/export?workspace=demo', json=dict(revision=0, format='xml'))
    assert response.status_code == 200 and 'attachment' in response.headers['content-disposition']
    changed = State.model_validate(original); changed.settings.name = 'Копия школы'
    result = client.post('/api/exchange/xml-preview?workspace=demo', json=dict(revision=0, xml=state_xml(changed).decode())).json()
    assert result['canApply'] and result['diff']['changed'] == 1
    assert client.get('/api/state?workspace=demo').json() == original
    applied = client.post('/api/editor/apply?workspace=demo', json=dict(revision=0, token=result['token']))
    assert applied.json()['settings']['name'] == 'Копия школы'
    assert client.post('/api/exchange/export?workspace=demo', json=dict(revision=0, format='xlsx')).status_code == 409
