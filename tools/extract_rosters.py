"""Read roster exports offline. Never execute scripts or fetch saved-page resources.

Run with Python containing openpyxl and beautifulsoup4; the output is plain JSON
consumed by import_rosters.py. Original source files are never changed.
"""
import argparse
from collections import Counter
from email import policy
from email.parser import BytesParser
import hashlib
import json
from pathlib import Path
import re
import unicodedata
from urllib.parse import urlparse


def clean(value):
    return ' '.join(unicodedata.normalize('NFC', str(value or '')).split())


def person_key(name, class_name):
    return clean(name).casefold(), clean(class_name).casefold()


def digest(value):
    return hashlib.sha256(value.encode('utf8')).hexdigest()[:20]


def read_mhtml(path):
    from bs4 import BeautifulSoup
    message = BytesParser(policy=policy.default).parsebytes(Path(path).read_bytes())
    part = next((p for p in message.walk() if p.get_content_type() == 'text/html'), None)
    if part is None:
        raise ValueError(f'{path}: нет HTML-части')
    raw = part.get_payload(decode=True)
    # Blink can omit the MIME charset; the saved document declares UTF-8.
    match = re.search(br'<meta[^>]+charset=["\s]*([\w-]+)', raw[:10000], re.I)
    charset = match.group(1).decode('ascii') if match else part.get_content_charset() or 'utf-8'
    soup = BeautifulSoup(raw.decode(charset, errors='strict'), 'html.parser')
    tables = soup.select('table.studentPredmetTable')
    if len(tables) != 1:
        raise ValueError(f'{path}: ожидалась одна таблица составов')
    table = tables[0]
    host = urlparse(message.get('Snapshot-Content-Location', '')).hostname
    if not host:
        raise ValueError(f'{path}: нет адреса школы для идентификаторов учеников')
    students, checks, skipped = {}, {}, 0
    for row in table.select('tr'):
        inputs = row.select('input[data-studentid]')
        if not inputs:
            continue
        cells = row.find_all('td', recursive=False)
        ids = {i['data-studentid'] for i in inputs}
        if len(ids) != 1 or len(cells) < 2:
            raise ValueError('Неоднозначная строка ученика')
        name, class_name = clean(cells[0].get_text(' ', strip=True)), clean(cells[1].get_text())
        # EduPage stores class attendance placeholders alongside real students.
        if re.fullmatch(r'Все\s+\d+\s*Присутствуют', name):
            if any(i.has_attr('checked') for i in inputs):
                raise ValueError('Служебная строка отмечена как участник группы')
            skipped += 1
            continue
        if not name or not re.fullmatch(r'\d+[A-Za-zА-Яа-я]', class_name):
            raise ValueError(f'Не распознана строка ученика: {class_name}')
        external = f'{host}:{next(iter(ids))}'
        if external in students:
            raise ValueError('Повторная строка ученика')
        students[external] = dict(key=external, externalId=external, name=name,
                                  className=class_name, active='studentIsOutTr' not in row.get('class', []))
        checks[external] = {}
        for i in inputs:
            plan = i.get('data-planid', '')
            if plan.startswith('na:'):
                continue
            cell = i.find_parent('td')
            checked = i.has_attr('checked')
            if checked != ('cellSelected' in cell.get('class', [])):
                raise ValueError('Атрибут checkbox и выделение ячейки расходятся')
            if plan in checks[external]:
                raise ValueError('Повторное назначение ученика в колонке')
            checks[external][plan] = checked
    courses = []
    for header in table.select('tr.header th'):
        group = header.select_one('[title^="Группа:"]')
        if group is None:
            continue
        plan = header.select_one('[data-planid]')['data-planid']
        subject = group.find_next_sibling('div')
        teacher = header.find_all('div', recursive=False)[1]
        classes = header.select_one('[title^="Класс:"]')['title'].split(':', 1)[1].strip().split('+')
        if any(plan not in values for values in checks.values()):
            raise ValueError(f'{path}: неполная колонка {plan}')
        members = [key for key, values in checks.items() if values[plan]]
        footer = table.select_one(f'td.footerCell[data-planid="{plan}"]')
        if footer is None or clean(footer.get_text()) != str(len(members)):
            raise ValueError(f'{path}: итог колонки {plan} не совпадает с отметками')
        courses.append(dict(name=clean(group.get_text()), subject=clean(subject.get_text()),
                            teacher=clean(teacher.get('title') or teacher.get_text()),
                            classNames=sorted(classes), members=sorted(members),
                            externalId=f'{host}:course:{plan}', sources=[Path(path).name],
                            sourceType='mhtml', kind='elective'))
    return dict(students=list(students.values()), courses=courses, skipped=skipped)


def read_xlsx(path):
    import openpyxl
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    students, courses = [], []
    try:
        for sheet in workbook:
            rows = list(sheet.iter_rows(values_only=True))
            if len(rows) < 5 or clean(rows[0][1]) != 'Класс':
                raise ValueError(f'{sheet.title}: неизвестный формат заголовка')
            columns = []
            for column in range(2, len(rows[1])):
                name = clean(rows[1][column])
                if not name:
                    continue
                course = dict(name=name, subject=clean(rows[2][column]), teacher=clean(rows[3][column]),
                              classNames=[], members=[], sources=[f'{Path(path).name} · лист {sheet.title}'],
                              sourceType='xlsx', kind='math' if name in ['Math A', 'Math B', 'Math C'] else 'core')
                courses.append(course)
                columns.append((column, course))
            seen = set()
            totals = None
            for number, row in enumerate(rows[4:], 5):
                name, class_name = clean(row[0]), clean(row[1])
                if not name and class_name == '#':
                    if totals is not None:
                        raise ValueError('Повторная итоговая строка')
                    totals = row
                    continue
                if not name and not class_name:
                    if any(v not in (None, '') for v in row[2:]):
                        raise ValueError(f'{sheet.title}!{number}: значения без имени')
                    continue
                if not name or not re.fullmatch(r'\d+[A-Za-zА-Яа-я]', class_name):
                    raise ValueError(f'{sheet.title}!{number}: неизвестная строка')
                key = person_key(name, class_name)
                if key in seen:
                    raise ValueError(f'{sheet.title}: повтор имени внутри класса')
                seen.add(key)
                sid = 'xlsx:' + digest('|'.join(key))
                students.append(dict(key=sid, externalId=None, name=name, className=class_name, active=True))
                for column, course in columns:
                    mark = row[column]
                    if mark not in (None, '', 0, 1, '0', '1'):
                        raise ValueError(f'{sheet.title}!{number}: неизвестная отметка {mark!r}')
                    if class_name not in course['classNames']:
                        course['classNames'].append(class_name)
                    if mark in (1, '1'):
                        course['members'].append(sid)
            for column, course in columns:
                course['rosterStatus'] = 'complete'
                if totals is not None and clean(totals[column]) != str(len(course['members'])):
                    course['rosterStatus'] = 'partial'
                    course['note'] = f'В Excel отмечено {len(course["members"])} учеников, итоговая строка показывает {clean(totals[column])}. Требуется сверка состава.'
    finally:
        workbook.close()
    return dict(students=students, courses=courses)


def extract(workbook, snapshots):
    people, names, html_courses, stats = {}, {}, {}, []
    for path in snapshots:
        data = read_mhtml(path)
        stats.append(dict(file=Path(path).name, students=len(data['students']), courses=len(data['courses']),
                          inactive=sum(not s['active'] for s in data['students']), skipped=data['skipped']))
        for student in data['students']:
            key = student['key']
            if key in people and people[key] != student:
                raise ValueError('Снимки расходятся в данных ученика; нужна сверка')
            nk = person_key(student['name'], student['className'])
            if nk in names and names[nk] != key:
                raise ValueError('Тёзки с разными идентификаторами; требуется явное сопоставление')
            people[key], names[nk] = student, key
        for course in data['courses']:
            key = course['externalId']
            if key in html_courses:
                prior = html_courses[key]
                if any(prior[k] != course[k] for k in ['name', 'subject', 'teacher', 'classNames', 'members']):
                    raise ValueError(f'Снимки расходятся в составе {course["name"]}')
                prior['sources'] = sorted(set(prior['sources'] + course['sources']))
            else:
                html_courses[key] = course
    core = read_xlsx(workbook)
    aliases = {}
    for student in core['students']:
        nk = person_key(student['name'], student['className'])
        key = names.get(nk, student['key'])
        aliases[student['key']] = key
        if key not in people:
            people[key], names[nk] = student, key
        elif not people[key]['active']:
            raise ValueError('Ученик из Excel отмечен выбывшим в снимке: нужна сверка')
    for c in core['courses']:
        c['members'] = sorted(aliases[s] for s in c['members'])
        c['classNames'].sort()
    stats.insert(0, dict(file=Path(workbook).name, students=len(core['students']), courses=len(core['courses']), inactive=0, skipped=0))
    # The user designates the workbook as the authority for core / math.
    courses = list(core['courses'])
    comparisons = []
    for c in html_courses.values():
        matches = [x for x in core['courses'] if (x['name'], x['subject'], x['classNames']) == (c['name'], c['subject'], c['classNames'])]
        if matches:
            if len(matches) != 1:
                raise ValueError('Неоднозначная core-группа')
            authoritative = matches[0]
            comparisons.append(dict(name=c['name'], sources=c['sources'], workbookCount=len(authoritative['members']),
                                    snapshotCount=len(c['members']), sameMembers=c['members'] == authoritative['members']))
            continue
        if c['subject'] == 'Math' or c['subject'] == 'Further Math':
            raise ValueError('Несопоставленная основная группа в снимке')
        courses.append(c)
    for c in courses:
        if any(people[s]['className'] not in c['classNames'] for s in c['members']):
            raise ValueError('Ученик за пределами классов курса')
    return dict(students=list(people.values()), courses=courses, sources=stats, comparisons=comparisons)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--core', required=True)
    parser.add_argument('--electives', nargs='+', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    data = extract(args.core, args.electives)
    Path(args.output).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf8')
    print(json.dumps(dict(students=len(data['students']), courses=len(data['courses']), sources=data['sources'], comparisons=data['comparisons']), ensure_ascii=False))
